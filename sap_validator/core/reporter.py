"""
Report generator for SAP Migration Post-Load Validation.
Produces:
  1. Excel workbook  (summary + per-field mismatch tabs)
  2. HTML report     (standalone, browser-friendly)
"""

from __future__ import annotations
import os
import json
from datetime import datetime
from pathlib import Path
from typing import TYPE_CHECKING

if TYPE_CHECKING:
    from core.validator import ValidationResult


# ── Excel Report ─────────────────────────────────────────────────────────────
def generate_excel_report(result: "ValidationResult", output_path: str) -> str:
    try:
        import openpyxl
        from openpyxl.styles import (
            PatternFill, Font, Alignment, Border, Side
        )
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise ImportError("pip install openpyxl")

    wb = openpyxl.Workbook()

    GREEN  = "FF00AA44"
    RED    = "FFCC2200"
    YELLOW = "FFFF9900"
    GREY   = "FF555555"
    LIGHT_GREEN = "FFE6F4EA"
    LIGHT_RED   = "FFFCE8E6"
    LIGHT_GREY  = "FFF5F5F5"
    WHITE  = "FFFFFFFF"

    def hdr_fill(hex_color):
        return PatternFill("solid", fgColor=hex_color)

    def thin_border():
        s = Side(style="thin", color="FFCCCCCC")
        return Border(left=s, right=s, top=s, bottom=s)

    # ── Sheet 1: Executive Summary ───────────────────────────────────────────
    ws = wb.active
    ws.title = "Summary"
    ws.sheet_view.showGridLines = False

    # Title
    ws.merge_cells("A1:F1")
    c = ws["A1"]
    c.value = "SAP Material Master — Post-Load Validation Report"
    c.font = Font(bold=True, size=16, color=WHITE)
    c.fill = hdr_fill("FF1B3A57")
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    # Run metadata
    meta = [
        ("Run Date", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Source File", result.source_file),
        ("Target File", result.target_file),
        ("Overall Status", result.overall_status),
    ]
    for i, (k, v) in enumerate(meta, start=3):
        ws.cell(i, 1, k).font = Font(bold=True)
        ws.cell(i, 2, v)
        status_color = GREEN if result.overall_status == "PASS" else RED
        if k == "Overall Status":
            ws.cell(i, 2).font = Font(bold=True,
                color=GREEN if v == "PASS" else RED)

    # Record counts
    ws["A8"] = "Record Counts"
    ws["A8"].font = Font(bold=True, size=12)
    rows = [
        ("Source Records",  result.total_source_records),
        ("Target Records",  result.total_target_records),
        ("Matched (both)",  result.records_matched),
        ("Only in Source",  result.records_only_in_source),
        ("Only in Target",  result.records_only_in_target),
    ]
    for i, (k, v) in enumerate(rows, start=9):
        ws.cell(i, 1, k)
        ws.cell(i, 2, v)
        if k in ("Only in Source", "Only in Target") and v > 0:
            ws.cell(i, 2).font = Font(color="FFCC2200", bold=True)

    # Field summary table
    ws["A16"] = "Field-Level Validation Summary"
    ws["A16"].font = Font(bold=True, size=12)
    headers = ["Source Field", "Target Field", "Records", "Matched",
               "Mismatched", "Miss-Source", "Miss-Target", "Match %", "Status"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(17, col, h)
        cell.fill = hdr_fill("FF1B3A57")
        cell.font = Font(bold=True, color=WHITE, size=10)
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border()

    for row_i, fr in enumerate(result.field_results, start=18):
        data = [fr.field_source, fr.field_target, fr.total_records,
                fr.matched, fr.mismatched, fr.missing_in_source,
                fr.missing_in_target, f"{fr.match_pct}%", fr.status]
        bg = LIGHT_GREEN if fr.status == "PASS" else LIGHT_RED
        for col, val in enumerate(data, 1):
            cell = ws.cell(row_i, col, val)
            cell.fill = PatternFill("solid", fgColor=bg)
            cell.border = thin_border()
            cell.alignment = Alignment(horizontal="center" if col > 2 else "left")
            if col == len(data):  # status column
                cell.font = Font(
                    bold=True,
                    color=GREEN if fr.status == "PASS" else RED
                )

    col_widths = [22, 28, 10, 10, 12, 12, 12, 10, 10]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # ── Sheet 2+: Mismatch detail per failing field ───────────────────────────
    failing = [fr for fr in result.field_results if fr.status == "FAIL"]
    for fr in failing:
        if not fr.mismatch_details:
            continue
        safe_name = fr.field_source[:28].replace("/", "_").replace("\\", "_")
        ws2 = wb.create_sheet(title=f"FAIL_{safe_name}")
        ws2.sheet_view.showGridLines = False

        ws2.merge_cells("A1:D1")
        c2 = ws2["A1"]
        c2.value = f"Mismatches: {fr.field_source} → {fr.field_target}"
        c2.font = Font(bold=True, size=13, color=WHITE)
        c2.fill = hdr_fill("FFCC2200")
        c2.alignment = Alignment(horizontal="center", vertical="center")
        ws2.row_dimensions[1].height = 28

        detail_hdrs = ["Material Number", "Source Value", "Target Value", "Issue"]
        for col, h in enumerate(detail_hdrs, 1):
            cell = ws2.cell(3, col, h)
            cell.fill = hdr_fill("FF333333")
            cell.font = Font(bold=True, color=WHITE)
            cell.border = thin_border()

        for ri, rec in enumerate(fr.mismatch_details, start=4):
            row_vals = [
                rec.get("material", ""),
                rec.get("source_value", ""),
                rec.get("target_value", ""),
                rec.get("issue", ""),
            ]
            bg_alt = LIGHT_RED if ri % 2 == 0 else LIGHT_GREY
            for ci, v in enumerate(row_vals, 1):
                cell = ws2.cell(ri, ci, v)
                cell.fill = PatternFill("solid", fgColor=bg_alt)
                cell.border = thin_border()

        for col, w in zip("ABCD", [22, 28, 28, 32]):
            ws2.column_dimensions[col].width = w

    wb.save(output_path)
    return output_path


# ── HTML Report ───────────────────────────────────────────────────────────────
def generate_html_report(result: "ValidationResult", output_path: str) -> str:
    ss = result.summary_stats
    run_ts = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    status_badge = (
        '<span style="background:#00AA44;color:#fff;padding:4px 16px;'
        'border-radius:4px;font-weight:700;">PASS</span>'
        if result.overall_status == "PASS" else
        '<span style="background:#CC2200;color:#fff;padding:4px 16px;'
        'border-radius:4px;font-weight:700;">FAIL</span>'
    )

    field_rows = ""
    for fr in result.field_results:
        color = "#E6F4EA" if fr.status == "PASS" else "#FCE8E6"
        badge = (
            '<span style="color:#00AA44;font-weight:700;">PASS</span>'
            if fr.status == "PASS" else
            '<span style="color:#CC2200;font-weight:700;">FAIL</span>'
        )
        field_rows += f"""
        <tr style="background:{color}">
          <td>{fr.field_source}</td>
          <td>{fr.field_target}</td>
          <td>{fr.total_records}</td>
          <td>{fr.matched}</td>
          <td>{fr.mismatched}</td>
          <td>{fr.missing_in_source}</td>
          <td>{fr.missing_in_target}</td>
          <td><b>{fr.match_pct}%</b></td>
          <td>{badge}</td>
        </tr>"""

    mismatch_sections = ""
    for fr in result.field_results:
        if not fr.mismatch_details:
            continue
        detail_rows = "".join(
            f"<tr><td>{r['material']}</td><td>{r['source_value']}</td>"
            f"<td>{r['target_value']}</td><td>{r['issue']}</td></tr>"
            for r in fr.mismatch_details
        )
        mismatch_sections += f"""
        <h3 style="margin-top:32px;color:#CC2200">
          ⚠ {fr.field_source} → {fr.field_target}
          <small style="font-size:0.7em;color:#888">({len(fr.mismatch_details)} issues)</small>
        </h3>
        <table>
          <thead>
            <tr>
              <th>Material</th><th>Source Value</th>
              <th>Target Value</th><th>Issue</th>
            </tr>
          </thead>
          <tbody>{detail_rows}</tbody>
        </table>"""

    html = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<title>SAP Migration Validation Report</title>
<style>
  body{{font-family:'Segoe UI',Arial,sans-serif;margin:0;background:#F3F6F9;color:#222}}
  .wrap{{max-width:1200px;margin:0 auto;padding:32px 24px}}
  h1{{color:#1B3A57;margin-bottom:4px}}
  .meta{{color:#666;font-size:0.9em;margin-bottom:24px}}
  .cards{{display:grid;grid-template-columns:repeat(auto-fit,minmax(160px,1fr));gap:16px;margin:24px 0}}
  .card{{background:#fff;border-radius:8px;padding:20px;box-shadow:0 1px 4px rgba(0,0,0,.08);text-align:center}}
  .card .num{{font-size:2em;font-weight:700;color:#1B3A57}}
  .card .lbl{{font-size:0.8em;color:#888;margin-top:4px}}
  .card.warn .num{{color:#CC2200}}
  table{{width:100%;border-collapse:collapse;background:#fff;border-radius:8px;
         overflow:hidden;box-shadow:0 1px 4px rgba(0,0,0,.08);margin-bottom:32px}}
  th{{background:#1B3A57;color:#fff;padding:10px 12px;text-align:left;font-size:0.85em}}
  td{{padding:9px 12px;border-bottom:1px solid #eee;font-size:0.9em}}
  tr:last-child td{{border-bottom:none}}
  h2{{color:#1B3A57;margin-top:40px}}
  h3{{color:#333}}
</style>
</head>
<body>
<div class="wrap">
  <h1>SAP Material Master — Post-Load Validation Report</h1>
  <div class="meta">
    Run: {run_ts} &nbsp;|&nbsp;
    Source: <code>{result.source_file}</code> &nbsp;|&nbsp;
    Target: <code>{result.target_file}</code> &nbsp;|&nbsp;
    Status: {status_badge}
  </div>

  <div class="cards">
    <div class="card">
      <div class="num">{result.total_source_records}</div>
      <div class="lbl">Source Records</div>
    </div>
    <div class="card">
      <div class="num">{result.total_target_records}</div>
      <div class="lbl">Target Records</div>
    </div>
    <div class="card">
      <div class="num">{result.records_matched}</div>
      <div class="lbl">Keys Matched</div>
    </div>
    <div class="card {'warn' if result.records_only_in_source else ''}">
      <div class="num">{result.records_only_in_source}</div>
      <div class="lbl">Only in Source</div>
    </div>
    <div class="card {'warn' if result.records_only_in_target else ''}">
      <div class="num">{result.records_only_in_target}</div>
      <div class="lbl">Only in Target</div>
    </div>
    <div class="card">
      <div class="num">{ss['fields_passed']}/{ss['total_fields_validated']}</div>
      <div class="lbl">Fields Passed</div>
    </div>
    <div class="card {'warn' if ss['fields_failed'] else ''}">
      <div class="num">{ss['fields_failed']}</div>
      <div class="lbl">Fields Failed</div>
    </div>
    <div class="card">
      <div class="num">{ss['pass_rate_pct']}%</div>
      <div class="lbl">Field Pass Rate</div>
    </div>
  </div>

  <h2>Field-Level Summary</h2>
  <table>
    <thead>
      <tr>
        <th>Source Field</th><th>Target Field</th><th>Records</th>
        <th>Matched</th><th>Mismatched</th>
        <th>Miss-Source</th><th>Miss-Target</th>
        <th>Match %</th><th>Status</th>
      </tr>
    </thead>
    <tbody>{field_rows}</tbody>
  </table>

  {'<h2>Mismatch Details</h2>' + mismatch_sections if mismatch_sections else ''}
</div>
</body>
</html>"""

    Path(output_path).write_text(html, encoding="utf-8")
    return output_path
